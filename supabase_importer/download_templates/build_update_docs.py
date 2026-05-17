from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
GRAY = RGBColor(85, 85, 85)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
BORDER = "9AA7B7"


def dxa(inches: float) -> int:
    return int(round(inches * 1440))


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_font(paragraph, size=11, color=BLACK, bold=False, italic=False):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=180, bottom=120, end=180):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    widths_dxa = [dxa(w) if isinstance(w, float) else int(w) for w in widths]
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx >= len(widths_dxa):
                continue
            cell.width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = tr_pr.find(qn("w:tblHeader"))
        if tbl_header is None:
            tbl_header = OxmlElement("w:tblHeader")
            tr_pr.append(tbl_header)
        tbl_header.set(qn("w:val"), "true")


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def make_header_row(table, labels):
    for cell, label in zip(table.rows[0].cells, labels):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, size=10.5, color=NAVY, bold=True)
        set_cell_shading(cell, HEADER_FILL)


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    make_header_row(table, headers)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(text))
            set_run_font(run, size=font_size, color=BLACK)
    add_spacer(doc, 5)
    return table


def add_form_table(doc, rows, label_width=2.55, fill_width=3.95, blank_lines=1):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [label_width, fill_width])
    set_table_borders(table, color="B9C3D0", size="4")
    make_header_row(table, ["项目与填写要求", "生产商填写"])

    for row in rows:
        if len(row) == 2:
            label, requirement = row
            lines = blank_lines
        else:
            label, requirement, lines = row
        cells = table.add_row().cells

        left = cells[0]
        left.text = ""
        label_p = left.paragraphs[0]
        label_p.paragraph_format.space_after = Pt(3)
        label_run = label_p.add_run(label)
        set_run_font(label_run, size=10.2, color=NAVY, bold=True)
        req_p = left.add_paragraph()
        req_p.paragraph_format.space_after = Pt(0)
        req_p.paragraph_format.line_spacing = 1.15
        req_run = req_p.add_run(requirement)
        set_run_font(req_run, size=8.8, color=GRAY)

        right = cells[1]
        right.text = ""
        for idx in range(max(1, lines)):
            p = right.paragraphs[0] if idx == 0 else right.add_paragraph()
            p.paragraph_format.space_after = Pt(10 if lines > 1 else 0)
            p.add_run("")
    add_spacer(doc, 7)
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_table_borders(table, color="D0D7E2", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, size=9.5, color=BLACK)
    add_spacer(doc, 7)


def add_spacer(doc, points):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(points)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def add_mono_lines(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        set_run_font(r, size=9.5, color=BLACK, name="Consolas")
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    add_spacer(doc, 6)


def setup_document(title_for_header):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run(title_for_header)
    set_run_font(hr, size=9, color=GRAY, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("HeatEx Direct | 生产商资料更新模板")
    set_run_font(fr, size=8.5, color=GRAY)
    return doc


def add_title_block(doc, kicker, title, subtitle, rows):
    brand = doc.add_paragraph()
    brand.paragraph_format.space_after = Pt(1)
    brand_run = brand.add_run("HeatEx Direct")
    set_run_font(brand_run, size=16, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(kicker)
    set_run_font(r, size=10.5, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, size=24, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    set_run_font(r, size=12.5, color=GRAY)

    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1.45, 5.05])
    set_table_borders(table, color="D7DBE2", size="4")
    for idx, (label, value) in enumerate(rows):
        c0, c1 = table.rows[idx].cells
        set_cell_shading(c0, HEADER_FILL)
        for cell, text, bold in ((c0, label, True), (c1, value, False)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=9.5, color=NAVY if bold else BLACK, bold=bold)
    add_spacer(doc, 10)


INDUSTRY_ROWS = [
    ("1", "HVAC", "暖通空调"),
    ("2", "Petrochemical", "石油化工"),
    ("3", "Marine", "船舶/海工"),
    ("4", "Energy", "能源"),
    ("5", "Food Processing", "食品加工"),
    ("6", "Chemical Processing", "化工"),
    ("7", "Food and Beverage", "食品饮料"),
    ("8", "Pharmaceuticals", "制药"),
    ("9", "Power Generation", "发电/电力"),
]


def build_company_doc():
    doc = setup_document("Company Update Template")
    add_title_block(
        doc,
        "Supplier Submission Form",
        "公司信息更新表",
        "请生产商按表填写公司资料，并按指定文件夹打包附件。",
        [
            ("适用范围", "公司主体信息、公司级证书、客户 logo、公司图片、公司级文档"),
            ("语言要求", "请尽量同时提供中文和英文。若只提供中文，网站发布时的英文翻译可能无法完全准确。"),
            ("填写方式", "左侧为项目和要求，右侧留空给生产商填写"),
        ],
    )
    add_callout(
        doc,
        "填写原则",
        "请填写真实、可核验的信息。没有资料时留空，不要编造证书、客户、出口市场、工厂面积、员工人数或技术能力。请尽量提供中英文两版内容；如果只提供中文，网站发布时的英文翻译可能无法完全准确。",
    )

    add_heading(doc, "1. 公司基础信息", 1)
    add_form_table(
        doc,
        [
            ("公司英文目录名", "用于资料文件夹命名。请使用小写英文字母、数字和连字符，不要使用中文、空格或特殊符号。"),
            ("公司英文全称", "请填写官方注册名或官网展示名。"),
            ("公司中文全称", "请填写官方中文名称；没有中文名称可留空。"),
            ("公司一句话英文简介", "建议 1 句话，说明主营产品、行业和核心能力。请尽量提供英文原文。", 2),
            ("公司一句话中文简介", "建议 1 句话，与英文简介含义一致。", 2),
            ("公司英文详细介绍", "建议 3-5 个自然段；可介绍历史、产品范围、制造能力、行业经验和服务范围。请尽量提供英文原文。", 5),
            ("公司中文详细介绍", "建议 3-5 个自然段；与英文内容一致。", 5),
            ("公司优势（中英文）", "逐条填写真实优势，例如制造能力、认证资质、测试能力、服务网络、典型客户等。建议每条同时提供中文和英文。", 4),
            ("官网", "请填写完整网址，以 http:// 或 https:// 开头。"),
            ("公司视频链接", "如有 YouTube、官网视频或其他公开视频链接请填写；无则留空。"),
            ("社交媒体链接", "如 LinkedIn、YouTube 等；请注明平台名称和链接。", 2),
        ],
    )

    add_heading(doc, "2. 联系方式与经营信息", 1)
    add_form_table(
        doc,
        [
            ("公开联系邮箱", "可填写多个邮箱；请确认可以公开展示。", 2),
            ("公开联系电话", "建议包含国家区号；可填写多个号码。", 2),
            ("国家/地区", "请填写英文和中文名称，例如 China / 中国。"),
            ("城市", "请填写英文和中文名称。"),
            ("详细地址", "填写可公开展示的办公室、工厂或注册地址。", 2),
            ("成立年份", "只填写可核验年份；未知则留空。"),
            ("工厂面积", "请保留单位，例如 20000 m2；未知或无法核验则留空。"),
            ("员工人数", "可填写范围或约数，例如 200+；未知或无法核验则留空。"),
            ("适用行业", "从下方行业名称中选择，可多选。", 2),
            ("出口市场", "填写真实出口国家/地区；可写国家英文名、中文名或两字母代码。", 3),
            ("页面标题建议", "如需指定平台展示标题可填写；无则由平台整理。"),
            ("页面简介建议", "如需指定平台搜索/页面简介可填写；建议简洁真实。", 2),
        ],
    )

    add_heading(doc, "3. 行业选项", 1)
    add_table(doc, ["编号", "英文名称", "中文名称"], INDUSTRY_ROWS, [0.65, 2.65, 3.20], font_size=9.8)

    add_heading(doc, "4. 附件文件夹结构", 1)
    add_body(doc, "请把附件放入以下固定文件夹。即使某类资料暂时没有，也请保留空文件夹。")
    add_mono_lines(
        doc,
        [
            "提交资料包/",
            "  <公司英文目录名>/",
            "    company_certifications/",
            "    company_customers/",
            "    company_docs/",
            "    company_images/",
            "    product_certifications/",
            "    product_docs/",
            "    product_images/",
        ],
    )
    add_form_table(
        doc,
        [
            ("company_certifications 文件夹", "放公司级证书、体系认证、生产资质等。请在右侧列出文件名、证书名称、有效期/到期日和简短说明；证书需在提交资料当天仍然有效。", 4),
            ("company_customers 文件夹", "放真实合作客户 logo。请在右侧列出客户名称和文件名；不要提交未经确认的客户。", 4),
            ("company_docs 文件夹", "放公司简介、企业画册、公司级白皮书等。只接受 pdf、doc、docx，单文件小于 50MB。", 4),
            ("company_images 文件夹", "放工厂外观、车间、设备、检测、团队等公司展示图。图片需真实、清晰、可公开展示。", 4),
        ],
    )
    add_callout(
        doc,
        "文件命名要求",
        "建议文件名使用小写英文字母、数字和下划线，例如 iso_9001_certificate.jpg、factory_exterior.jpg。请不要使用中文、空格或特殊符号。文件名必须与实际提交文件完全一致。",
    )

    add_heading(doc, "5. 公司附件清单", 1)
    add_form_table(
        doc,
        [
            ("公司证书清单", "每行建议包含：证书名称、文件名、颁发机构、有效期/到期日、说明。证书需在提交资料当天仍然有效。", 6),
            ("客户 logo 清单", "每行建议包含：客户名称、文件名、是否可公开展示。", 5),
            ("公司文档清单", "每行建议包含：文档标题、文件名、语言、文件类型。只接受 pdf、doc、docx。", 5),
            ("公司图片清单", "每行建议包含：图片主题、文件名、展示说明。", 6),
        ],
    )

    add_heading(doc, "6. 提交前确认", 1)
    add_form_table(
        doc,
        [
            ("资料真实性确认", "确认所有公司信息、证书、客户 logo、图片和文档均真实且可公开使用。"),
            ("证书有效期确认", "确认所有公司级证书在提交资料当天仍然有效。"),
            ("文件夹确认", "确认 7 个固定文件夹均已保留，附件放入正确文件夹。"),
            ("文件名确认", "确认表中填写的文件名与实际提交文件完全一致。"),
            ("文档格式确认", "确认公司文档只包含 pdf、doc、docx，且单文件小于 50MB。"),
            ("空白项确认", "确认无法核实的信息已留空，没有用猜测内容填充。"),
        ],
    )

    doc.save(OUT_DIR / "company_information_update_template.docx")


def build_product_doc():
    doc = setup_document("Product Update Template")
    add_title_block(
        doc,
        "Supplier Submission Form",
        "产品信息更新表",
        "请生产商按表填写产品资料，并按指定文件夹打包产品附件。",
        [
            ("适用范围", "产品基础信息、产品优势、技术参数、详情内容、产品图片、产品证书、产品文档"),
            ("语言要求", "请尽量同时提供中文和英文。若只提供中文，网站发布时的英文翻译可能无法完全准确。"),
            ("填写方式", "左侧为项目和要求，右侧留空给生产商填写"),
        ],
    )
    add_callout(
        doc,
        "产品保留门槛",
        "每个产品至少需要一段真实介绍，并至少提供产品优势、技术参数、详情说明或真实产品文档中的一类。只有名称、页面标题或图片的占位产品不要提交。",
    )

    add_heading(doc, "1. 产品基础信息", 1)
    add_form_table(
        doc,
        [
            ("所属公司英文目录名", "必须与公司信息更新表中的公司英文目录名一致。"),
            ("产品英文目录名", "用于产品页面和文件整理。请使用小写英文字母、数字和连字符，不要使用中文、空格或特殊符号。"),
            ("产品显示顺序", "如有多个产品，请填写显示顺序，例如 1、2、3。"),
            ("产品英文名称", "使用官网、样本或产品资料中的正式名称。"),
            ("产品中文名称", "使用官网、样本或产品资料中的正式名称；无中文名可留空。"),
            ("产品一句话英文简介", "建议 1 句话，说明产品类型、主要应用或核心优势。请尽量提供英文原文。", 2),
            ("产品一句话中文简介", "建议 1 句话，与英文简介含义一致。", 2),
            ("产品英文详细介绍", "介绍产品结构、适用工况、行业应用、性能特点等。请尽量提供英文原文。", 5),
            ("产品中文详细介绍", "与英文内容一致。", 5),
            ("适用行业", "从下方行业名称中选择，可多选。", 2),
            ("产品视频链接", "如有产品视频请填写完整网址；无则留空。"),
        ],
    )

    add_heading(doc, "2. 产品内容信息", 1)
    add_form_table(
        doc,
        [
            ("产品优势（中英文）", "逐条填写真实优势，例如高效传热、耐压、易维护、材料选择、模块化等。建议每条同时提供中文和英文。", 6),
            ("技术参数（中英文）", "每行填写一个参数：参数名称、数值或范围、单位、备注。参数名称建议同时提供中英文；不要填写无法核实的参数。", 7),
            ("产品详情说明（中英文）", "可填写问答、结构说明、工作原理、制造方式、应用场景等；建议每条包含标题和内容，并同时提供中英文。", 7),
            ("页面标题建议", "如需指定平台展示标题可填写；无则由平台整理。"),
            ("页面简介建议", "如需指定平台搜索/页面简介可填写；建议简洁真实。", 2),
        ],
    )

    add_heading(doc, "3. 行业选项", 1)
    add_table(doc, ["编号", "英文名称", "中文名称"], INDUSTRY_ROWS, [0.65, 2.65, 3.20], font_size=9.8)

    add_heading(doc, "4. 产品附件文件夹结构", 1)
    add_body(doc, "产品附件与公司附件放在同一个公司资料包下。产品文档必须明确归属到某一个已填写产品，不能放入公司文档文件夹。")
    add_mono_lines(
        doc,
        [
            "提交资料包/",
            "  <公司英文目录名>/",
            "    company_certifications/",
            "    company_customers/",
            "    company_docs/",
            "    company_images/",
            "    product_certifications/",
            "    product_docs/",
            "    product_images/",
        ],
    )
    add_form_table(
        doc,
        [
            ("product_images 文件夹", "放产品图、结构图、应用图。只提交方图或横图，图片宽度必须大于或等于高度；竖图不要提交。", 4),
            ("product_certifications 文件夹", "放产品级证书、型式认证、产品资质等。请在右侧列出文件名、证书名称、有效期/到期日和适用产品；证书需在提交资料当天仍然有效。", 4),
            ("product_docs 文件夹", "放产品手册、规格书、样本、说明书。只接受 pdf、doc、docx，单文件小于 50MB；必须说明对应产品。", 4),
        ],
    )
    add_callout(
        doc,
        "归属规则",
        "产品手册、产品说明书、产品规格书和产品样本必须放入 product_docs，并在表中写明对应产品。如果文档无法明确对应到某个产品，请不要作为公司文档提交。",
    )

    add_heading(doc, "5. 产品附件清单", 1)
    add_form_table(
        doc,
        [
            ("产品图片清单", "每行建议包含：产品名称、图片主题、文件名、图片方向确认。图片必须为方图或横图。", 7),
            ("产品证书清单", "每行建议包含：适用产品、证书名称、文件名、颁发机构、有效期/到期日、说明。证书需在提交资料当天仍然有效。", 7),
            ("产品文档清单", "每行建议包含：对应产品、文档标题、文件名、语言、文件类型。只接受 pdf、doc、docx。", 7),
        ],
    )

    add_heading(doc, "6. 提交前确认", 1)
    add_form_table(
        doc,
        [
            ("产品内容确认", "确认每个产品都有真实介绍，并至少具备优势、参数、详情说明或产品文档之一。"),
            ("产品图片确认", "确认产品图片均为方图或横图，没有竖图。"),
            ("产品证书有效期确认", "确认所有产品级证书在提交资料当天仍然有效。"),
            ("产品文档归属确认", "确认每个产品文档都能明确对应到已填写产品。"),
            ("文件名确认", "确认表中填写的文件名与实际提交文件完全一致。"),
            ("文档格式确认", "确认产品文档只包含 pdf、doc、docx，且单文件小于 50MB。"),
            ("空白项确认", "确认无法核实的信息已留空，没有用猜测内容填充。"),
        ],
    )

    doc.save(OUT_DIR / "product_information_update_template.docx")


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_company_doc()
    build_product_doc()
